import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io

st.set_page_config(page_title="Trias Rooster Generator", page_icon="📅", layout="centered")

st.title("📅 Trias Rooster Generator")
st.markdown("Upload de CSV export uit Airtable en download het opgemaakte Excel rooster.")

uploaded_file = st.file_uploader("Kies je CSV bestand", type="csv")

if uploaded_file:
    st.success(f"✓ Bestand geladen: {uploaded_file.name}")

    if st.button("Genereer rooster", type="primary"):
        with st.spinner("Rooster wordt gegenereerd..."):

            LES_KLEUREN = [
                ('C6EFCE','375623'), ('BDD7EE','1F4E79'), ('E2D4F0','4B2E83'),
                ('FFF2CC','7F6000'), ('D9EAD3','274E13'), ('CFE2F3','0B5394'),
                ('EAD1DC','741B47'), ('F4CCCC','85200C'), ('DDEBF7','1F3864'),
            ]

            df = pd.read_csv(uploaded_file)
            df_les = df[df['Type'] == 'Les'].copy()
            df_les['Datum'] = pd.to_datetime(df_les['Datum'], dayfirst=True)
            df_les = df_les.sort_values('Datum')

            df_vak = df[df['Type'] == 'Vakantie'].copy()
            df_vak['Datum'] = pd.to_datetime(df_vak['Datum'], dayfirst=True)
            vakantie_datums = set(df_vak['Datum'].dt.date)
            vakantie_namen = {}
            for _, row in df_vak.iterrows():
                vakantie_namen[row['Datum'].date()] = row['Notitie'] if pd.notna(row['Notitie']) else 'Vakantie'

            dagen_volgorde = ['Maandag', 'Dinsdag', 'Woensdag', 'Donderdag', 'Vrijdag']
            scholen = [s for s in df_les['School'].dropna().unique()]

            wb = Workbook()
            wb.remove(wb.active)

            def maak_border(dik_onder=False):
                return Border(
                    left=Side(style='thin', color='BFBFBF'),
                    right=Side(style='thin', color='BFBFBF'),
                    top=Side(style='thin', color='BFBFBF'),
                    bottom=Side(style='medium' if dik_onder else 'thin', color='9E9E9E' if dik_onder else 'BFBFBF')
                )

            center = Alignment(horizontal='center', vertical='center', wrap_text=True)
            left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)

            for school in scholen:
                ws = wb.create_sheet(title=school[:31])
                df_school = df_les[df_les['School'] == school].copy()

                unieke_templates = df_school['Les templates'].dropna().unique()
                template_kleur = {t: LES_KLEUREN[i % len(LES_KLEUREN)] for i, t in enumerate(unieke_templates)}

                ws.merge_cells('A1:H1')
                ws['A1'] = f'Rooster {school} – Schooljaar 2025-2026'
                ws['A1'].font = Font(name='Arial', bold=True, size=13, color='FFFFFF')
                ws['A1'].fill = PatternFill('solid', start_color='1F4E79')
                ws['A1'].alignment = center
                ws.row_dimensions[1].height = 28

                headers = ['Lesweek','Week','Datum Ma-Vr','Maandag','Dinsdag','Woensdag','Donderdag','Vrijdag']
                for col, h in enumerate(headers, 1):
                    cell = ws.cell(row=2, column=col, value=h)
                    cell.font = Font(name='Arial', bold=True, size=10, color='FFFFFF')
                    cell.fill = PatternFill('solid', start_color='2E75B6')
                    cell.alignment = center
                    cell.border = maak_border()
                ws.row_dimensions[2].height = 18

                ws.column_dimensions['A'].width = 8
                ws.column_dimensions['B'].width = 7
                ws.column_dimensions['C'].width = 16
                for col in ['D','E','F','G','H']:
                    ws.column_dimensions[col].width = 26

                start_datum = datetime(2025, 9, 1)
                eind_datum = datetime(2026, 7, 18)
                huidige = start_datum
                while huidige.weekday() != 0:
                    huidige -= timedelta(days=1)

                lesweek = 1
                rij = 3

                while huidige <= eind_datum:
                    maandag = huidige
                    vrijdag = huidige + timedelta(days=4)
                    week_nr = maandag.isocalendar()[1]
                    datum_str = f"{maandag.strftime('%d %b')} - {vrijdag.strftime('%d %b')}"

                    week_datums = [maandag + timedelta(days=i) for i in range(5)]
                    is_vakantie = any(d.date() in vakantie_datums for d in week_datums)

                    if is_vakantie:
                        vak_naam = next((vakantie_namen[d.date()] for d in week_datums if d.date() in vakantie_namen), 'Vakantie')
                        for col, val in [(1,''),(2,week_nr),(3,datum_str)]:
                            cell = ws.cell(row=rij, column=col, value=val)
                            cell.fill = PatternFill('solid', start_color='FFE699')
                            cell.font = Font(name='Arial', size=8, color='7F4F00')
                            cell.alignment = center
                            cell.border = maak_border(dik_onder=True)
                        ws.merge_cells(start_row=rij, start_column=4, end_row=rij, end_column=8)
                        vak_cell = ws.cell(row=rij, column=4, value=vak_naam)
                        vak_cell.fill = PatternFill('solid', start_color='FFE699')
                        vak_cell.font = Font(name='Arial', bold=True, size=9, color='7F4F00')
                        vak_cell.alignment = center
                        vak_cell.border = maak_border(dik_onder=True)
                        for col in range(5, 9):
                            cell = ws.cell(row=rij, column=col)
                            cell.fill = PatternFill('solid', start_color='FFE699')
                            cell.border = maak_border(dik_onder=True)
                        ws.row_dimensions[rij].height = 18
                        rij += 1
                    else:
                        lessen_per_dag = []
                        for dag_i in range(5):
                            dag_datum = maandag + timedelta(days=dag_i)
                            lessen = df_school[df_school['Datum'].dt.date == dag_datum.date()]
                            lessen_per_dag.append(lessen)

                        max_lessen = max(len(l) for l in lessen_per_dag) if any(len(l) > 0 for l in lessen_per_dag) else 1

                        for sub_rij in range(max_lessen):
                            is_laatste = (sub_rij == max_lessen - 1)
                            rand = maak_border(dik_onder=is_laatste)

                            for col, val in [(1, lesweek if sub_rij==0 else ''), (2, week_nr if sub_rij==0 else ''), (3, datum_str if sub_rij==0 else '')]:
                                cell = ws.cell(row=rij, column=col, value=val)
                                cell.fill = PatternFill('solid', start_color='D6E4F0')
                                cell.font = Font(name='Arial', size=8)
                                cell.alignment = center
                                cell.border = rand

                            for dag_i, lessen in enumerate(lessen_per_dag):
                                col = dag_i + 4
                                if sub_rij < len(lessen):
                                    les = lessen.iloc[sub_rij]
                                    naam = str(les['Les templates']).replace(f' - {school}', '').strip()
                                    tijd = f"{les['Starttijd']}-{les['Eindtijd']}"
                                    docent = les['Docent'] if pd.notna(les['Docent']) else ''
                                    tekst = f"{naam}\n{tijd}  |  {docent}"
                                    gaat_door = les['Gaat door'] == 'checked'
                                    bg, fg = ('FCE4D6','833C00') if not gaat_door else template_kleur.get(les['Les templates'], ('E2EFDA','375623'))
                                    cell = ws.cell(row=rij, column=col, value=tekst)
                                    cell.fill = PatternFill('solid', start_color=bg)
                                    cell.font = Font(name='Arial', size=8, color=fg)
                                else:
                                    cell = ws.cell(row=rij, column=col, value='')
                                    cell.fill = PatternFill('solid', start_color='F9F9F9')
                                cell.alignment = left_wrap
                                cell.border = rand

                            ws.row_dimensions[rij].height = 30
                            rij += 1

                        lesweek += 1

                    huidige += timedelta(weeks=1)

                # Legenda
                rij += 1
                ws.cell(row=rij, column=1, value='Legenda:').font = Font(name='Arial', bold=True, size=9)
                rij += 1
                col_leg = 1
                for template, (bg, fg) in template_kleur.items():
                    naam_kort = template.replace(f' - {school}', '').strip()
                    cell = ws.cell(row=rij, column=col_leg, value=naam_kort)
                    cell.fill = PatternFill('solid', start_color=bg)
                    cell.font = Font(name='Arial', size=8, color=fg)
                    cell.border = maak_border()
                    cell.alignment = left_wrap
                    col_leg += 1
                    if col_leg > 4:
                        col_leg = 1
                        rij += 1

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

        st.success("✓ Rooster succesvol gegenereerd!")
        st.download_button(
            label="⬇️ Download Excel rooster",
            data=output,
            file_name="Trias_Rooster_2025-2026.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
